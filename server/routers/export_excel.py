"""
Excel导出功能
"""
from fastapi import APIRouter, Query, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from database import get_db
from deps import get_current_user
from models.sale import SaleOrder
from models.purchase import PurchaseOrder
from models.inventory import InventoryRecord
from models.finance import Receipt, Payment
from models.expense import Expense
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

router = APIRouter(prefix="/api/export", tags=["导出"])

THIN = Side(style='thin', color='999999')
THIN_BORDER = Border(left=THIN, top=THIN, right=THIN, bottom=THIN)
HEADER_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
HEADER_FONT = Font(color='FFFFFF', bold=True, size=11)
CENTER = Alignment(horizontal='center', vertical='center')


def set_header(ws, headers, col_widths):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + col)].width = width


def set_row(ws, row_idx, values, is_sum=False):
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col, value=val)
        cell.font = Font(size=10, bold=is_sum)
        cell.border = THIN_BORDER
        if isinstance(val, (int, float)):
            cell.alignment = Alignment(horizontal='right')
        else:
            cell.alignment = CENTER


def money(v):
    return round(float(v or 0), 2)


def make_wb():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


@router.get("/profit")
def export_profit(
    start_date: str = Query(None),
    end_date: str = Query(None),
    db: Session = Depends(get_db),
    _=Depends(get_current_user)
):
    from models.report import ProfitReport
    q = db.query(ProfitReport)
    if start_date:
        q = q.filter(ProfitReport.date >= start_date)
    if end_date:
        q = q.filter(ProfitReport.date <= end_date)
    rows = q.order_by(ProfitReport.date).all()

    wb = make_wb()
    ws = wb.create_sheet("利润报表")
    set_header(ws, ['日期', '销售金额', '成本金额', '利润', '利润率'], [14, 14, 14, 14, 12])
    total_sales = total_cost = 0
    for i, r in enumerate(rows, 2):
        s = money(r.sales_amount)
        c = money(r.cost_amount)
        p = money(r.profit)
        rate = r.gross_rate or 0
        total_sales += s
        total_cost += c
        set_row(ws, i, [r.date, s, c, p, f'{rate*100:.2f}%'])
    if rows:
        tp = total_sales - total_cost
        trate = tp / total_sales if total_sales else 0
        set_row(ws, len(rows) + 2, ['合计', total_sales, total_cost, tp, f'{trate*100:.2f}%'], is_sum=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"利润报表_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    from urllib.parse import quote
    encoded_filename = quote(filename)
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                            headers={'Content-Disposition': f'attachment; filename*=UTF-8\'\'{encoded_filename}'})
